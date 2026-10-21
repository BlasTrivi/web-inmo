#!/usr/bin/env python
"""
Regenera src/data/units.ts a partir de la lista de precios mensual (Excel).

Uso (desde la raiz del proyecto):
    python -I scripts/units-from-xlsx.py "ruta/a/Listas de precios - DD-MM-AAAA.xlsx"

Requiere:  pip install openpyxl
Lee cada hoja (President Tower, The Address, Boston, Marquis, Trinity III), detecta el bloque de
"Condiciones comerciales" y la tabla de unidades por el texto de los encabezados (no por posicion
de columna), saltea filas ocultas/vacias y escribe src/data/units.ts.
Hojas nuevas: agregarlas en SHEET_TO_SLUG. Revisar siempre el resumen que imprime al final.
"""
import datetime
import json
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl

SHEET_TO_SLUG = {
    'President Tower': 'president-tower',
    'The Address': 'the-address',
    'Boston': 'boston-842',
    'Marquis': 'marquis',
    'Trinity III': 'trinity-3',
}
# Proyectos cuya pagina tiene la seccion #tipologias (la tipologia de cada unidad enlaza ahi)
TIPOLOGIAS_LINK = {'president-tower', 'the-address', 'boston-842', 'trinity-3'}
OUT = Path(__file__).resolve().parent.parent / 'src' / 'data' / 'units.ts'

BAD = '�'  # el Excel viene con tildes rotas
# Correcciones conocidas
FIXES = [
    (BAD + 'ltimo', 'Último'), ('d' + BAD + 'a', 'día'), (r'\bhabil\b', 'hábil'), ('variaci' + BAD + 'n', 'variación'),
    ('posesi' + BAD + 'n', 'posesión'), ('seg' + BAD + 'n', 'según'), ('tipolog' + BAD + 'a', 'tipología'),
    ('D' + BAD + 'plex', 'Dúplex'), ('Balc' + BAD + 'n', 'Balcón'), ('ACTUALIZACI' + BAD + 'N', 'ACTUALIZACIÓN'),
    (r'\bM' + BAD + r'\b', 'm²'),
]


def fix(s):
    if not isinstance(s, str):
        return s
    for pat, rep in FIXES:
        s = re.sub(pat, rep, s)
    if BAD in s:
        print('  AVISO: texto con caracter roto sin corregir:', repr(s))
    return re.sub(r'\s+', ' ', s).strip()


def norm(s):
    """Encabezado normalizado: sin tildes, minusculas, m2/m² -> m, sin signos."""
    s = unicodedata.normalize('NFD', str(s or ''))
    s = ''.join(c for c in s if unicodedata.category(c) != 'Mn').lower()
    s = s.replace(BAD, '').replace('²', '').replace('*', '')
    s = re.sub(r'\bm2\b', 'm', s)
    return re.sub(r'\s+', ' ', s).strip()


def num(v, nd=2):
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        v = round(float(v), nd)
        return int(v) if v == int(v) else v
    return None


def first_right(ws, r):
    cells = [c.value for c in ws[r] if c.value is not None]
    return cells[1] if len(cells) > 1 else None


def parse_sheet(ws, slug):
    hidden = {r for r, d in ws.row_dimensions.items() if d.hidden}
    cond = {}
    header_row = None
    for r in range(1, ws.max_row + 1):
        a = norm(ws.cell(r, 1).value)
        if a == 'piso':
            header_row = r
            break
        v = first_right(ws, r)
        if v is None:
            continue
        if a.startswith('precio cochera'):
            cond['cocheraAdicional'] = num(v)
        elif a.startswith('precio baulera'):
            cond['baulera'] = num(v)
        elif a.startswith('fecha de actualizacion'):
            cond['_updated'] = v.date().isoformat() if isinstance(v, datetime.datetime) else str(v)
        elif a.startswith('vigencia'):
            cond['vigencia'] = fix(v)
        elif a.startswith('forma de pago'):
            cond['formaPago'] = fix(v)
        elif a.startswith('condiciones generales'):
            cond['generales'] = fix(v)
    if header_row is None:
        raise SystemExit(f'Hoja {ws.title}: no encontre la fila de encabezados (celda "Piso" en col. A)')

    # campo -> columna (primera aparicion; las columnas de PRECIO de cochera/baulera se ignoran)
    cols = {}
    for c in range(1, ws.max_column + 1):
        h = norm(ws.cell(header_row, c).value)
        if not h:
            continue
        key = None
        if h == 'piso': key = 'piso'
        elif h.startswith('unid'): key = 'unidad'
        elif h.startswith('tipolog'): key = 'tipologia'
        elif h.startswith('observ'): key = 'obs'
        elif h == 'm cubiertos': key = 'cubiertos'
        elif h.startswith('m balcon'): key = 'balcon'
        elif h == 'm propios': key = 'total'
        elif h == 'cocheras': key = 'cocheras'
        elif h == 'baulera': key = 'baulera'
        elif h in ('precio contado', 'valor'): key = 'precio'
        if key and key not in cols:
            cols[key] = c
    for need in ('piso', 'tipologia', 'total', 'precio'):
        if need not in cols:
            raise SystemExit(f'Hoja {ws.title}: falta la columna "{need}" (encabezados: {cols})')

    units = []
    for r in range(header_row + 1, ws.max_row + 1):
        if r in hidden:
            continue

        def get(k):
            return ws.cell(r, cols[k]).value if k in cols else None

        piso_raw = get('piso')
        if not (isinstance(piso_raw, str) and norm(piso_raw).startswith('piso')):
            continue  # fila vacia, nota al pie, etc.
        piso = re.sub(r'(?i)^piso\s*', '', fix(piso_raw)).replace(' / ', ' · ')
        tip = fix(get('tipologia') or '')
        tip = re.sub(r'(?i)\bdpto\b', 'departamento', tip)
        tip = re.sub(r'(?i)\bdorm\b', 'dormitorios', tip)
        tip = re.sub(r'(?i)\+ serv\b', '+ servicio', tip)
        tip = re.sub(r'(?i)\bdormitorios\b', 'dormitorios', tip)
        tip = tip[:1].upper() + tip[1:]
        m = re.search(r'(\d)\s*dorm', tip, re.I)
        obs = fix(get('obs') or '') or None
        letter = fix(get('unidad')) if get('unidad') else None
        u = {'unidad': letter if letter else f'Piso {piso}'}
        if letter:
            u['piso'] = piso
        u['tipologia'] = tip
        if m:
            u['dormitorios'] = int(m.group(1))
        for k, f in (('cubiertos', 'm2Cubiertos'), ('balcon', 'm2Balcon'), ('total', 'm2')):
            if num(get(k)) is not None:
                u[f] = num(get(k))
        if num(get('cocheras')) is not None:
            u['cocheras'] = int(num(get('cocheras')))
        if num(get('baulera')) is not None:
            u['baulera'] = int(num(get('baulera')))
        u['estado'] = 'Disponible'
        if obs:
            if re.search(r'alquil|modelo', obs, re.I):
                u['estado'] = obs[:1].upper() + obs[1:]
            else:
                u['nota'] = obs
        p = num(get('precio'), 0)
        if p is not None:
            u['precio'] = int(p)
        units.append(u)
    updated = cond.pop('_updated', None)
    return {'updated': updated, 'conditions': cond, 'tipologiasLink': slug in TIPOLOGIAS_LINK, 'units': units}


HEADER = """/**
 * Unidades disponibles por proyecto. ARCHIVO GENERADO por `scripts/units-from-xlsx.py` desde la lista de
 * precios mensual (Excel de INMO). No editar a mano salvo correcciones puntuales.
 *
 * COMO ACTUALIZAR CADA MES:
 *   1. Guardar el Excel nuevo (una hoja por proyecto, mismo formato).
 *   2. python -I scripts/units-from-xlsx.py "ruta/al/excel.xlsx"   (necesita `pip install openpyxl`)
 *   3. Revisar el resumen que imprime y el diff de src/data/units.ts.
 * Proyectos sin hoja (Trinity 2, Laguna Club) no figuran acá: la sección muestra el aviso "consultá con un asesor".
 *
 * Los PRECIOS y las condiciones comerciales (forma de pago, cochera/baulera) están cargados pero NO se
 * publican mientras `SHOW_PRICES` sea false (src/config/site.ts).
 */

export interface AvailableUnit {
  /** Identificador visible: "C", "A+B", "Piso 5" */
  unidad: string;
  piso?: string;
  /** "3 dormitorios + servicio", "Semipiso 3 dormitorios + servicio" */
  tipologia?: string;
  dormitorios?: number;
  /** m² cubiertos, balcón interno y totales (propios = cubiertos + balcón + cocheras + baulera) */
  m2Cubiertos?: number;
  m2Balcon?: number;
  m2?: number;
  /** Cocheras / bauleras incluidas en el precio */
  cocheras?: number;
  baulera?: number;
  /** "Disponible", "Alquilada", "Unidad modelo..." */
  estado?: string;
  /** Aclaración de la planilla ("Dúplex", "Piso completo") */
  nota?: string;
  /** Precio de contado en USD (número). Solo se muestra con SHOW_PRICES = true */
  precio?: number;
  /** URL del plano (opcional) */
  plano?: string;
}

export interface UnitConditions {
  /** USD */
  cocheraAdicional?: number;
  baulera?: number;
  vigencia?: string;
  formaPago?: string;
  generales?: string;
}

export interface ProjectUnits {
  /** Fecha de la lista (AAAA-MM-DD) */
  updated: string;
  conditions: UnitConditions;
  /** La página del proyecto tiene sección #tipologias: la tipología de cada unidad enlaza ahí */
  tipologiasLink: boolean;
  units: AvailableUnit[];
}

"""


def main():
    if len(sys.argv) != 2:
        raise SystemExit(__doc__)
    wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
    out = {}
    for ws in wb:
        slug = SHEET_TO_SLUG.get(ws.title.strip())
        if not slug:
            print(f'Hoja ignorada (sin mapeo en SHEET_TO_SLUG): {ws.title}')
            continue
        out[slug] = parse_sheet(ws, slug)
    dump = lambda o: json.dumps(o, ensure_ascii=False)
    lines = [HEADER, 'export const UNITS: Record<string, ProjectUnits> = {']
    for slug, d in out.items():
        lines.append(f"  '{slug}': {{")
        lines.append(f"    updated: '{d['updated']}',")
        lines.append(f"    conditions: {dump(d['conditions'])},")
        lines.append(f"    tipologiasLink: {'true' if d['tipologiasLink'] else 'false'},")
        lines.append('    units: [')
        for u in d['units']:
            lines.append(f'      {dump(u)},')
        lines.append('    ],')
        lines.append('  },')
    lines.append('};')
    OUT.write_text('\n'.join(lines) + '\n', encoding='utf-8')
    print(f'Escrito {OUT}')
    for slug, d in out.items():
        print(f"  {slug}: {len(d['units'])} unidades, actualizada {d['updated']}, condiciones: {sorted(d['conditions'])}")


if __name__ == '__main__':
    main()
